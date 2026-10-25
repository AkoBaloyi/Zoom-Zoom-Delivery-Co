#!/usr/bin/env python3
"""Convert one of this project's Markdown documents to .docx, and optionally to PDF.

Written because the briefs and milestone documents are easier to read and to hand to a
collaborator as Word or PDF than as raw Markdown. Handles the constructs the project's docs
actually use: ATX headings, pipe tables, fenced code blocks, ordered and unordered lists,
horizontal rules, and inline bold, italic, inline code and links.

    python docs/tooling/md_to_docx.py docs/UI-TUTORIAL-BRIEF.md
    python docs/tooling/md_to_docx.py docs/UI-TUTORIAL-BRIEF.md --pdf
    python docs/tooling/md_to_docx.py docs/M2-INDIVIDUAL-AKO.md --out build/M2.docx --pdf

PDF export drives the installed Microsoft Word through COM, so it only works on Windows with
Word present. Without --pdf the script has no Windows-specific dependency.

Mermaid fenced blocks are emitted as captioned monospace source rather than rendered diagrams;
Word cannot render Mermaid. Export those from the Markdown preview and insert them manually.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    sys.exit("python-docx is required: python -m pip install python-docx")

MONO = "Consolas"
CODE_FILL = "F4F4F4"
RULE_COLOUR = "BFBFBF"

# Inline: `code`, **bold**, *italic*, [text](url). Code first so its contents are left alone.
INLINE = re.compile(
    r"(?P<code>`[^`]+`)"
    r"|(?P<bold>\*\*[^*]+\*\*)"
    r"|(?P<italic>(?<!\*)\*(?!\*)[^*]+\*(?!\*))"
    r"|(?P<link>\[[^\]]+\]\([^)]+\))"
)


def shade(paragraph, fill: str) -> None:
    """Give a paragraph a solid background fill."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shd)


def horizontal_rule(doc) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), RULE_COLOUR)
    pbdr.append(bottom)
    p._p.get_or_add_pPr().append(pbdr)


def add_inline(paragraph, text: str, base_bold: bool = False) -> None:
    """Add text to a paragraph, honouring inline Markdown."""
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos : m.start()])
            r.bold = base_bold
        if m.group("code"):
            r = paragraph.add_run(m.group("code")[1:-1])
            r.font.name = MONO
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        elif m.group("bold"):
            paragraph.add_run(m.group("bold")[2:-2]).bold = True
        elif m.group("italic"):
            paragraph.add_run(m.group("italic")[1:-1]).italic = True
        elif m.group("link"):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", m.group("link")).groups()
            r = paragraph.add_run(label)
            r.bold = base_bold
            r.underline = True
            r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            if not label.strip().endswith(url.strip()):
                tail = paragraph.add_run(f" ({url})")
                tail.font.size = Pt(8)
                tail.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        r.bold = base_bold


def split_row(line: str) -> list[str]:
    inner = line.strip()
    if inner.startswith("|"):
        inner = inner[1:]
    if inner.endswith("|"):
        inner = inner[:-1]
    return [c.strip() for c in inner.split("|")]


def is_separator(line: str) -> bool:
    return bool(re.fullmatch(r"\s*\|?[\s:\-|]+\|?\s*", line)) and "-" in line


def add_table(doc, rows: list[list[str]]) -> None:
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j in range(width):
            cell = cells[j]
            cell.paragraphs[0].text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after = Pt(2)
            text = row[j] if j < len(row) else ""
            add_inline(para, text, base_bold=(i == 0))
            for run in para.runs:
                if run.font.size is None:
                    run.font.size = Pt(9)
            if i == 0:
                shade(para, "E8E8E8")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_code_block(doc, lines: list[str], lang: str) -> None:
    if lang.lower() == "mermaid":
        cap = doc.add_paragraph()
        r = cap.add_run("[Mermaid diagram — render from the Markdown preview and insert here]")
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x80, 0x30, 0x30)
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.25)
        r = p.add_run(line if line else " ")
        r.font.name = MONO
        r.font.size = Pt(8.5)
        shade(p, CODE_FILL)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def convert(md_path: Path, out_path: Path) -> Path:
    text = md_path.read_text(encoding="utf-8", errors="replace")
    lines = text.splitlines()

    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.9)
        section.top_margin = section.bottom_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            lang = stripped[3:].strip()
            body: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                body.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, body, lang)
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and is_separator(lines[i + 1]):
            rows = [split_row(line)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        if re.fullmatch(r"\s*([-*_])\s*(\1\s*){2,}", line):
            horizontal_rule(doc)
            i += 1
            continue

        h = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if h:
            level = len(h.group(1))
            heading = doc.add_heading("", level=min(level, 4))
            add_inline(heading, h.group(2))
            for run in heading.runs:
                run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
            i += 1
            continue

        b = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
        if b:
            depth = min(len(b.group(1)) // 2, 2)
            style = "List Bullet" if depth == 0 else f"List Bullet {depth + 1}"
            try:
                p = doc.add_paragraph(style=style)
            except KeyError:
                p = doc.add_paragraph(style="List Bullet")
            add_inline(p, b.group(2))
            i += 1
            continue

        n = re.match(r"^(\s*)\d+[.)]\s+(.*)$", line)
        if n:
            depth = min(len(n.group(1)) // 2, 2)
            style = "List Number" if depth == 0 else f"List Number {depth + 1}"
            try:
                p = doc.add_paragraph(style=style)
            except KeyError:
                p = doc.add_paragraph(style="List Number")
            add_inline(p, n.group(2))
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        # Gather a wrapped paragraph: Markdown hard-wraps, Word should reflow.
        buf = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            s = nxt.strip()
            if (
                not s
                or s.startswith(("|", "```", "#", "- ", "* ", "+ "))
                or re.match(r"^\s*\d+[.)]\s", nxt)
                or re.fullmatch(r"\s*([-*_])\s*(\1\s*){2,}", nxt)
            ):
                break
            buf.append(s)
            i += 1
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline(p, " ".join(buf))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def to_pdf(docx_path: Path) -> Path:
    """Export the .docx to PDF using the installed Word. Windows + Word only."""
    try:
        import win32com.client  # type: ignore
    except ImportError:
        try:
            import comtypes.client as cc  # type: ignore
        except ImportError:
            sys.exit(
                "PDF export needs pywin32 or comtypes:\n"
                "  python -m pip install pywin32\n"
                f"The .docx is written and usable: {docx_path}"
            )
        word = cc.CreateObject("Word.Application")
    else:
        word = win32com.client.Dispatch("Word.Application")

    pdf_path = docx_path.with_suffix(".pdf")
    word.Visible = False
    doc = None
    try:
        doc = word.Documents.Open(str(docx_path.resolve()))
        doc.SaveAs(str(pdf_path.resolve()), FileFormat=17)  # 17 = wdFormatPDF
    finally:
        if doc is not None:
            doc.Close(False)
        word.Quit()
    return pdf_path


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("markdown", type=Path, help="source .md file")
    ap.add_argument("--out", type=Path, help="output .docx (default: alongside the source)")
    ap.add_argument("--pdf", action="store_true", help="also export a PDF via Word")
    a = ap.parse_args()

    if not a.markdown.exists():
        sys.exit(f"not found: {a.markdown}")

    out = a.out or a.markdown.with_suffix(".docx")
    docx_path = convert(a.markdown, out)
    print(f"wrote {docx_path}  ({docx_path.stat().st_size:,} bytes)")

    if a.pdf:
        pdf_path = to_pdf(docx_path)
        print(f"wrote {pdf_path}  ({pdf_path.stat().st_size:,} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
